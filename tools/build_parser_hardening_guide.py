from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("mydocs/LLM-Wiki-Parser-Production-Hardening-Guide.docx")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "666666"


def set_font(run, name="Microsoft YaHei", size=11, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(header), size=10, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), size=9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.style = "Code Block"
    for line in text.splitlines():
        run = p.add_run(line + "\n")
        set_font(run, name="Consolas", size=9)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
    code_style = styles.add_style("Code Block", 1)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(9)
    code_style.paragraph_format.left_indent = Inches(0.2)
    code_style.paragraph_format.right_indent = Inches(0.2)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.0
    p_pr = code_style._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("LLM Wiki 技术说明"), size=9, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("Parser 生产化加固说明"), size=24, bold=True, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_font(subtitle.add_run("改动内容、设计动机、前后差异、模块关系与使用指南"), size=13, color=MUTED)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(16)
    set_font(meta.add_run("对应 Spec：2026-07-19_23-21_parsers-production-hardening.md\n"), size=9.5, color=MUTED)
    set_font(meta.add_run("实现提交：10b396b  |  评审结果：PASS  |  测试：162 passed"), size=9.5, color=MUTED)

    doc.add_heading("1. 执行摘要", level=1)
    p = doc.add_paragraph()
    set_font(p.add_run("核心结论："), bold=True, color=DARK_BLUE)
    p.add_run("这次工作不是简单增加文件格式，而是把偏工具型的 parser 模块升级为可接入真实上传业务的生产解析流水线。")
    bullet(doc, "建立统一 ParseResult、稳定错误码和资源预算。")
    bullet(doc, "上传、Service、队列、Worker、Parser 和数据库形成完整状态链路。")
    bullet(doc, "显式选择解析引擎，不再静默回退或把未知格式当文本处理。")
    bullet(doc, "只有真正成功的内容才能进入分块、Embedding 和 content_chunks。")
    bullet(doc, "DOCX 图片得到安全持久化，Markdown 引用改写为稳定 MinIO URI。")

    doc.add_heading("2. 改造前后总览", level=1)
    doc.add_heading("2.1 改造前", level=2)
    add_code(doc, "上传文件 -> 保存 MinIO -> RQ worker -> 按扩展名找 parser\n          -> 返回纯文本 -> 分块 -> 向量化 -> 入库")
    doc.add_heading("2.2 改造后", level=2)
    add_code(doc, "上传文件 + parser_engine\n  -> 限制读取大小\n  -> 校验格式 / 文件大小 / 引擎可用性\n  -> 保存原文件和解析配置\n  -> RQ 异步任务\n  -> 严格选择 parser + 解析超时\n  -> ParseResult(content, images, metadata, warnings, error_code)\n  -> 图片持久化与 Markdown 路径改写\n  -> 分块 / Embedding / content_chunks\n  -> 保存状态与诊断信息")
    add_table(doc, ["维度", "以前", "现在"], [
        ("返回契约", "纯文本或异常", "结构化 ParseResult"),
        ("未知格式", "按 TXT 兜底，可能产生乱码", "unsupported_type，停止处理"),
        ("显式引擎失败", "可能回退 builtin", "engine_unavailable，不允许回退"),
        ("空内容", "可能标记 processed", "failed + empty_content"),
        ("图片", "容易丢失", "保存 MinIO 并改写正文引用"),
        ("诊断", "主要依赖错误文本", "稳定错误码 + metadata + warnings"),
    ], [1700, 3600, 4060])

    doc.add_heading("3. 具体修改与原因", level=1)
    doc.add_heading("3.1 统一解析结果契约", level=2)
    p = doc.add_paragraph("新增 backend/app/parsers/result.py，定义 ParseResult、ParseErrorCode、ParseLimits 和 ParseDispatchError。")
    p.add_run("目的：").bold = True
    p.add_run("让 API、worker、数据库和客户端共享同一套成功/失败语义，不再依赖易变的错误字符串。")
    add_table(doc, ["错误码", "含义", "调用方建议"], [
        ("unsupported_type", "不支持的文件类型", "提示更换格式"),
        ("parse_failed", "解析器执行失败", "记录日志并允许重试"),
        ("timeout", "解析超时", "建议换引擎或拆分文件"),
        ("too_large", "文件、正文或图片超限", "提示压缩或拆分"),
        ("empty_content", "没有可用正文", "检查扫描件或文件内容"),
        ("engine_unavailable", "引擎不存在、未安装或不支持格式", "换引擎或安装依赖"),
        ("unsafe_url", "URL 不安全", "拒绝请求"),
    ], [2200, 3400, 3760])

    doc.add_heading("3.2 严格派发与资源预算", level=2)
    p = doc.add_paragraph("dispatch.py 由宽松派发改为严格派发：未知类型不再走 TextParser，显式引擎不再回退 builtin，并统一检查输入、输出和图片预算。")
    add_table(doc, ["限制", "默认值", "保护目标"], [
        ("原文件大小", "50 MB", "HTTP 与内存"),
        ("解析正文长度", "5,000,000 字符", "内存、分块与向量成本"),
        ("派生图片总量", "20 MB", "内存与对象存储"),
        ("单次解析超时", "300 秒", "Worker 执行时间"),
        ("RQ Job 超时", "600 秒", "队列整体任务"),
    ], [3000, 2300, 4060])

    doc.add_heading("3.3 上传入口与 Service 前置校验", level=2)
    doc.add_paragraph("上传接口新增 multipart 字段 parser_engine，入口以 max+1 方式读取，Service 在写 MinIO 和创建任务之前验证格式、大小、引擎存在性、依赖可用性及格式兼容性。无效请求直接返回 HTTP 400 和机器可读 error_code。")
    doc.add_heading("3.4 Worker 状态机修正", level=2)
    doc.add_paragraph("Worker 改用 parse_document() 获取完整 ParseResult。parser 失败、超时、图片保存失败或空内容都进入 failed；只有完成分块、Embedding 和入库后才进入 processed 并设置 processed_at。失败内容不会污染 content_chunks。")
    add_code(doc, "pending -> processing -> processed\n                      \\-> failed(error_code, error_message, parse_metadata)")

    doc.add_heading("3.5 DOCX 图片、表格与 PDF 诊断", level=2)
    bullet(doc, "DOCX 同时提取段落、GFM Markdown 表格和内嵌图片。")
    bullet(doc, "图片以内容哈希生成确定性对象 key，重试可以幂等覆盖。")
    bullet(doc, "危险文件名会被清洗，Markdown 图片路径会改写为 minio:// URI。")
    bullet(doc, "DOCX 表格修复为连续 GFM 行，避免渲染器把每行识别为独立段落。")
    bullet(doc, "PDF 增加 is_scanned 元数据；扫描件当前只检测，不执行 OCR。")

    doc.add_heading("3.6 数据库字段与迁移", level=2)
    add_table(doc, ["字段", "作用", "典型值"], [
        ("parser_engine", "记录指定/实际解析引擎", "builtin"),
        ("parse_error_code", "机器可读失败原因", "timeout"),
        ("parse_metadata", "warnings、格式信息、图片对象等", "JSONB"),
    ], [2600, 4300, 2460])
    doc.add_paragraph("Alembic 002_parser_hardening.py 同时实现 upgrade 和 downgrade，支持添加或回滚这三列。")

    doc.add_heading("4. 模块关系与数据流", level=1)
    add_code(doc, "routers/routes.py\n  接收 file + parser_engine\n        |\nservices/document_service.py\n  校验请求，保存原文件与 Document\n        |\nworkers/queue.py -> RQ\n        |\nworkers/parse_document.py\n  控制状态、超时和全流程编排\n        |\nparsers/dispatch.py -> registry.py -> 具体 Parser\n        |\nparsers/result.py (统一结果契约)\n        |\nparser_asset_service.py (图片保存与引用改写)\n        |\nchunker + embedding -> content_chunks\n        |\ndocuments (状态、错误、metadata)")
    add_table(doc, ["模块", "核心职责", "不负责什么"], [
        ("Router", "HTTP 参数与有限读取", "不选择具体 parser"),
        ("Document Service", "业务校验、原文件和任务创建", "不执行耗时解析"),
        ("Registry", "维护 engine + file_type 映射", "不编排业务状态"),
        ("Dispatch", "严格选择 parser、统一结果和限制", "不写数据库"),
        ("Worker", "异步编排、超时、状态与后续入库", "不实现格式解析算法"),
        ("Asset Service", "派生图片安全持久化", "不解析文档正文"),
    ], [2100, 3900, 3360])

    doc.add_heading("5. 支持范围与引擎关系", level=1)
    doc.add_paragraph("Registry 的底层能力范围大于当前产品上传白名单。这样可以保留实验性能力，但在完成生产验证之前不对上传 API 开放。")
    add_table(doc, ["引擎", "生产使用方式", "依赖/限制"], [
        ("builtin", "默认；处理 txt、md、pdf、docx、xlsx、csv、pptx", "项目内置 parser"),
        ("markitdown", "可选；适合多种 Office/PDF 格式", "需安装 Microsoft MarkItDown 及格式依赖"),
        ("opendataloader", "可选；仅 PDF", "需 Java 11+ 和 opendataloader-pdf"),
    ], [2100, 3900, 3360])
    p = doc.add_paragraph()
    set_font(p.add_run("重要："), bold=True, color=DARK_BLUE)
    p.add_run("用户显式选择高级引擎时，如果依赖不可用或格式不受支持，系统直接失败，不会偷偷切回 builtin。")

    doc.add_heading("6. 使用指南", level=1)
    doc.add_heading("6.1 上传文档", level=2)
    add_code(doc, 'curl -X POST "http://localhost:8000/api/v1/kb/1/documents" \\\n  -H "Authorization: Bearer <token>" \\\n  -F "file=@report.docx" \\\n  -F "parser_engine=builtin"')
    doc.add_paragraph("MarkItDown：将 parser_engine 改为 markitdown。OpenDataLoader PDF：将其改为 opendataloader。")
    doc.add_heading("6.2 上传响应", level=2)
    add_code(doc, '{\n  "data": {\n    "doc_id": "f37c...",\n    "status": "pending",\n    "original_filename": "report.pdf",\n    "parser_engine": "builtin"\n  }\n}')
    doc.add_heading("6.3 查询状态", level=2)
    add_code(doc, 'curl "http://localhost:8000/api/v1/kb/1/documents/f37c.../status" \\\n  -H "Authorization: Bearer <token>"')
    doc.add_heading("6.4 客户端处理原则", level=2)
    numbered(doc, "先看 status：pending/processing 表示继续轮询，processed 表示完成，failed 表示停止。")
    numbered(doc, "失败时优先根据 error_code 决定提示和恢复动作。")
    numbered(doc, "error_message 只用于展示或日志，不要通过字符串匹配控制业务。")
    numbered(doc, "warnings 不等于失败；例如 partial_empty_pages 表示部分页面为空但仍可能有有效内容。")
    numbered(doc, "parse_metadata 可用于诊断引擎、格式和派生图片，不应假设其扩展字段永远固定。")

    doc.add_heading("7. 验证结果与剩余边界", level=1)
    add_table(doc, ["验证项", "结果"], [
        ("Parser、Service、Asset、Worker、Migration 测试", "162 passed"),
        ("触及文件 Ruff", "PASS"),
        ("FastAPI OpenAPI multipart 定义", "PASS"),
        ("Alembic 001 -> 002 离线升级", "PASS"),
        ("Alembic 002 -> 001 离线回滚", "PASS"),
        ("三轴 Review", "PASS，无 Blocking Issues"),
    ], [6500, 2860])
    doc.add_heading("7.1 已知边界", level=2)
    bullet(doc, "asyncio.to_thread 超时不能强杀底层线程；RQ 600 秒只是外层任务兜底。")
    bullet(doc, "本机未提供真实 PostgreSQL、Redis、MinIO 集成环境，当前证据以单元测试和离线迁移为主。")
    bullet(doc, "MinIO 网络中断可能留下部分派生对象；当前重试可幂等覆盖，但尚无孤立对象清理任务。")
    bullet(doc, "扫描 PDF 只检测，不包含 OCR。")
    doc.add_heading("7.2 建议后续顺序", level=2)
    numbered(doc, "建立 PostgreSQL、Redis、MinIO 真实基础设施集成测试。")
    numbered(doc, "将解析任务升级为可强制终止的进程级隔离。")
    numbered(doc, "补充解析指标、告警和孤立派生对象清理。")

    doc.add_heading("8. 文件改动索引", level=1)
    add_table(doc, ["类别", "关键文件"], [
        ("结果契约", "backend/app/parsers/result.py"),
        ("派发与注册", "backend/app/parsers/dispatch.py；registry.py；__init__.py"),
        ("格式解析", "docx2_parser.py；pdf_parser.py"),
        ("HTTP 与业务", "routers/routes.py；services/document_service.py"),
        ("异步执行", "workers/queue.py；workers/parse_document.py"),
        ("图片资产", "services/parser_asset_service.py；minio_service.py"),
        ("数据库", "models/routes.py；alembic/versions/002_parser_hardening.py"),
        ("配置与异常", "config/settings.py；config/exceptions.py；main.py"),
        ("验证", "test_parse_document_worker.py 等 7 组新增/增强测试"),
    ], [2300, 7060])

    doc.core_properties.title = "Parser 生产化加固说明"
    doc.core_properties.subject = "LLM Wiki parser production hardening"
    doc.core_properties.author = "LLM Wiki Project"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
