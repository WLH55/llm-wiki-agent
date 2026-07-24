"""PPTX 解析器契约测试。"""

from importlib import import_module
from io import BytesIO

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.parsers.core.base import BaseParser
from app.parsers.core.document import Document


def _pptx_parser_class():
    try:
        module = import_module("app.parsers.implementations.pptx")
    except ModuleNotFoundError:
        pytest.fail("缺少计划模块：app.parsers.implementations.pptx")
    return module.PptxParser


def _presentation_bytes() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[5])
    first.shapes.title.text = "项目概览"
    textbox = first.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1))
    textbox.text = "第一阶段完成"
    image_bytes = BytesIO()
    Image.new("RGB", (2, 2), color="red").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    first.shapes.add_picture(image_bytes, Inches(1), Inches(3), width=Inches(1))
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "下一步"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_pptx_parser_uses_base_parser_contract():
    assert issubclass(_pptx_parser_class(), BaseParser)


def test_parses_slide_text_and_embedded_media():
    document = _pptx_parser_class()(file_name="demo.pptx").parse(_presentation_bytes())
    assert isinstance(document, Document)
    assert "## Slide 1" in document.content
    assert "项目概览" in document.content
    assert "第一阶段完成" in document.content
    assert "## Slide 2" in document.content
    assert "下一步" in document.content
    assert "## 媒体" in document.content
    assert len(document.images) == 1
    image_path = next(iter(document.images))
    assert image_path.startswith("images/")
    assert f"![{image_path.removeprefix('images/')}]({image_path})" in document.content
    assert document.metadata == {
        "format": "pptx",
        "slide_count": 2,
        "media_count": 1,
    }


def test_extract_pptx_media_returns_base64_mapping():
    try:
        module = import_module("app.parsers.utils.pptx_media")
    except ModuleNotFoundError:
        pytest.fail("缺少计划模块：app.parsers.utils.pptx_media")
    images = module.extract_pptx_media(_presentation_bytes())
    assert list(images) == ["images/image1.png"]
    assert isinstance(images["images/image1.png"], str)
    assert images["images/image1.png"]


def test_legacy_ppt_is_rejected_without_external_conversion():
    module = import_module("app.parsers.implementations.pptx")
    legacy_ppt = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32
    with pytest.raises(ValueError, match="legacy_ppt_not_supported"):
        module.normalize_pptx_bytes(legacy_ppt, "ppt")


def test_registry_only_advertises_pptx():
    parser_class = _pptx_parser_class()
    from app.parsers import registry
    assert registry.get_parser_class("pptx") is parser_class
    with pytest.raises(KeyError):
        registry.get_parser_class("ppt")
