"""Excel parser suite contract tests."""

import zipfile
from importlib import import_module
from io import BytesIO

import openpyxl
import pytest

import app.parsers.excel_parser as excel_parser_module
import app.parsers.xlsx_repair as xlsx_repair_module
from app.parsers import registry
from app.parsers.base import BaseParser
from app.parsers.document import Document
from app.parsers.excel_convert import detect_excel_format, normalize_excel_bytes
from app.parsers.excel_parser import ExcelParser
from app.parsers.xlsx_merge import fill_merged_cells_xlsx
from app.parsers.xlsx_repair import repair_xlsx_bytes


@pytest.mark.parametrize(
    ("module_name", "symbols"),
    [
        ("app.parsers.excel_parser", ("ExcelParser",)),
        (
            "app.parsers.excel_convert",
            ("detect_excel_format", "convert_excel_to_xlsx_bytes", "normalize_excel_bytes"),
        ),
        ("app.parsers.xlsx_merge", ("fill_merged_cells_xlsx",)),
        ("app.parsers.xlsx_repair", ("repair_xlsx_bytes",)),
    ],
)
def test_excel_suite_exposes_planned_api(module_name, symbols):
    try:
        module = import_module(module_name)
    except ModuleNotFoundError:
        pytest.fail(f"planned module is missing: {module_name}")

    for symbol in symbols:
        assert hasattr(module, symbol), f"{module_name} must expose {symbol}"


def _workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "hello"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xlsx_with_phantom_shared_strings() -> bytes:
    files = {}
    with zipfile.ZipFile(BytesIO(_workbook_bytes()), "r") as source:
        for info in source.infolist():
            files[info.filename] = source.read(info.filename)

    content_types = files["[Content_Types].xml"].decode("utf-8")
    override = (
        '<Override PartName="/xl/sharedStrings.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.'
        'spreadsheetml.sharedStrings+xml"/>'
    )
    files["[Content_Types].xml"] = content_types.replace(
        "</Types>", override + "</Types>"
    ).encode("utf-8")

    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in files.items():
            target.writestr(name, data)
    return output.getvalue()


def test_excel_parser_uses_base_parser_contract():
    assert issubclass(ExcelParser, BaseParser)


def test_detects_xlsx_xls_and_unknown_bytes():
    assert detect_excel_format(_workbook_bytes()) == "xlsx"
    assert detect_excel_format(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32) == "xls"
    assert detect_excel_format(b"not a spreadsheet") is None


def test_fill_merged_cells_propagates_master_value():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "group"
    sheet.merge_cells("A1:B1")
    output = BytesIO()
    workbook.save(output)

    filled = fill_merged_cells_xlsx(output.getvalue())
    result = openpyxl.load_workbook(BytesIO(filled), data_only=True)

    assert result.active["A1"].value == "group"
    assert result.active["B1"].value == "group"
    assert not result.active.merged_cells.ranges


def test_repair_removes_phantom_shared_strings_reference():
    repaired = repair_xlsx_bytes(_xlsx_with_phantom_shared_strings())

    assert repaired is not None
    workbook = openpyxl.load_workbook(BytesIO(repaired), data_only=True)
    assert workbook.active["A1"].value == "hello"


def test_repair_returns_none_for_clean_workbook():
    assert repair_xlsx_bytes(_workbook_bytes()) is None


def _multi_sheet_workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    people = workbook.active
    people.title = "People"
    people.append(["Name", "Age", "Note"])
    people.append(["Alice", 30, "A|B"])
    people.append(["Bob", None, "line1\nline2"])

    flags = workbook.create_sheet("Flags")
    flags.append(["Enabled", "Count"])
    flags.append([False, 0])

    workbook.create_sheet("Empty")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parses_multiple_sheets_as_markdown_tables():
    document = ExcelParser(file_name="book.xlsx").parse(_multi_sheet_workbook_bytes())

    assert isinstance(document, Document)
    assert "## People" in document.content
    assert "| Name | Age | Note |" in document.content
    assert "| Alice | 30 | A\\|B |" in document.content
    assert "| Bob |  | line1<br>line2 |" in document.content
    assert "## Flags" in document.content
    assert "| FALSE | 0 |" in document.content
    assert "## Empty" not in document.content
    assert document.metadata == {
        "format": "xlsx",
        "sheet_count": 3,
        "parsed_sheet_count": 2,
    }


def test_invalid_excel_returns_error_metadata():
    document = ExcelParser(file_name="broken.xlsx").parse(b"not an excel file")

    assert document.content == ""
    assert document.metadata["error"].startswith("open_failed:")


def test_registry_routes_xlsx_to_excel_parser():
    assert registry.get_parser_class("xlsx") is ExcelParser


@pytest.mark.parametrize("file_type", ["xls", "xlsb", "ods", "et"])
def test_registry_does_not_advertise_formats_requiring_libreoffice(file_type):
    with pytest.raises(KeyError):
        registry.get_parser_class(file_type)


def test_external_conversion_is_disabled_in_default_normalization(monkeypatch):
    def unexpected_conversion(*args, **kwargs):
        raise AssertionError("default upload parsing must not start LibreOffice")

    monkeypatch.setattr(
        "app.parsers.excel_convert.convert_excel_to_xlsx_bytes",
        unexpected_conversion,
    )
    legacy_xls = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32

    with pytest.raises(ValueError, match="disabled"):
        normalize_excel_bytes(legacy_xls)


def test_sparse_worksheet_dimensions_are_rejected_before_iteration():
    validator = getattr(excel_parser_module, "_validate_sheet_dimensions", None)
    assert validator is not None, "Excel parser must expose an internal dimension guard"

    with pytest.raises(ValueError, match="worksheet_too_large"):
        validator("Sparse", max_row=1_048_576, max_column=16_384)


def test_archive_guard_rejects_extreme_compression_ratio():
    validator = getattr(xlsx_repair_module, "validate_xlsx_archive", None)
    assert validator is not None, "XLSX helpers must validate ZIP resource limits"

    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * 2_000_000)

    with pytest.raises(ValueError, match="compression_ratio"):
        validator(output.getvalue())
