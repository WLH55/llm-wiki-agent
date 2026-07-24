"""Word 系列 parser 测试。

- Docx2Parser：用 python-docx 在内存中构造一个 .docx bytes（含段落 + 表格），解析后断言
- DocParser：antiword/catdoc 通常未装，验证 fallback 行为（返回 metadata.error）
"""

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from docx.shared import Inches
from PIL import Image

from app.parsers.core.document import Document
from app.parsers.implementations.doc import DocParser
from app.parsers.implementations.docx import Docx2Parser


def _build_docx_bytes(paragraphs, table_rows=None) -> bytes:
    """用 python-docx 在内存中构造 .docx bytes。"""
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        rows = len(table_rows)
        cols = len(table_rows[0]) if table_rows else 0
        table = doc.add_table(rows=rows, cols=cols)
        for r, row_data in enumerate(table_rows):
            for c, val in enumerate(row_data):
                table.cell(r, c).text = str(val)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocx2Parser:
    def test_parses_paragraphs(self):
        payload = _build_docx_bytes(["第一段", "第二段"])
        parser = Docx2Parser(file_name="a.docx")
        doc = parser.parse(payload)
        assert isinstance(doc, Document)
        assert "第一段" in doc.content
        assert "第二段" in doc.content
        # 段落之间用双换行分隔
        assert doc.content.index("第一段") < doc.content.index("第二段")

    def test_parses_table(self):
        payload = _build_docx_bytes(
            paragraphs=["标题"],
            table_rows=[["姓名", "年龄"], ["张三", "25"]],
        )
        parser = Docx2Parser(file_name="t.docx")
        doc = parser.parse(payload)
        # 表格行被转成 markdown 风格（cell | cell）
        assert "姓名 | 年龄" in doc.content
        assert "张三 | 25" in doc.content
        assert (
            "| 姓名 | 年龄 |\n| --- | --- |\n| 张三 | 25 |" in doc.content
        )

    def test_extracts_images_with_markdown_reference(self):
        docx = DocxDocument()
        image_buffer = BytesIO()
        Image.new("RGB", (2, 2), color="blue").save(image_buffer, format="PNG")
        image_buffer.seek(0)
        docx.add_picture(image_buffer, width=Inches(1))
        payload = BytesIO()
        docx.save(payload)
        result = Docx2Parser(file_name="image.docx").parse(payload.getvalue())
        assert len(result.images) == 1
        image_path = next(iter(result.images))
        assert image_path.startswith("images/")
        assert f"]({image_path})" in result.content
        assert result.metadata["image_count"] == 1

    def test_empty_paragraphs_skipped(self):
        payload = _build_docx_bytes(["", "  ", "有效"])
        parser = Docx2Parser()
        doc = parser.parse(payload)
        assert doc.content.strip() == "有效"

    def test_invalid_bytes_returns_error_metadata(self):
        parser = Docx2Parser()
        doc = parser.parse(b"not a real docx")
        assert doc.content == ""
        assert "error" in doc.metadata
        assert "open_failed" in doc.metadata["error"]

    def test_empty_docx_returns_empty_content(self):
        payload = _build_docx_bytes([])
        parser = Docx2Parser()
        doc = parser.parse(payload)
        assert doc.content == ""
        assert doc.is_valid() is False
        assert doc.metadata == {"format": "docx", "image_count": 0}


class TestDocParser:
    """DocParser 依赖外部命令 antiword/catdoc。CI 环境 / 学习环境通常没装。

    本测试聚焦 fallback 语义：找不到工具时返回 metadata.error。
    如果环境碰巧装了，对应测试自然通过。
    """

    def test_returns_document_instance(self):
        parser = DocParser(file_name="a.doc")
        doc = parser.parse(b"fake doc bytes")
        assert isinstance(doc, Document)

    def test_no_tool_returns_error_metadata(self):
        import shutil

        if shutil.which("antiword") or shutil.which("catdoc"):
            pytest.skip("antiword/catdoc 已安装，跳过 fallback 验证")
        parser = DocParser(file_name="a.doc")
        doc = parser.parse(b"fake doc bytes")
        assert doc.content == ""
        assert doc.metadata.get("error") == "no_doc_tool"
        assert doc.metadata.get("tried") == ["antiword", "catdoc"]
