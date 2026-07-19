"""XLSX 合并单元格归一化工具。"""

import zipfile
from io import BytesIO

from openpyxl import load_workbook

from app.parsers.xlsx_repair import validate_xlsx_archive


def fill_merged_cells_xlsx(content: bytes) -> bytes:
    """取消合并区域，并将主单元格值复制到区域内所有单元格。"""
    if not zipfile.is_zipfile(BytesIO(content)):
        return content
    validate_xlsx_archive(content)
    from app.parsers.excel_parser import _validate_sheet_dimensions
    workbook = load_workbook(BytesIO(content), data_only=True)
    try:
        changed = False
        for sheet in workbook.worksheets:
            _validate_sheet_dimensions(sheet.title, sheet.max_row, sheet.max_column)
            for merged_range in list(sheet.merged_cells.ranges):
                master = sheet.cell(merged_range.min_row, merged_range.min_col).value
                sheet.unmerge_cells(str(merged_range))
                for row in range(merged_range.min_row, merged_range.max_row + 1):
                    for column in range(merged_range.min_col, merged_range.max_col + 1):
                        sheet.cell(row, column).value = master
                changed = True
        if not changed:
            return content
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
