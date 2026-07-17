"""Excel workbook parser with Markdown table output."""

import logging
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.parsers.base import BaseParser
from app.parsers.document import Document
from app.parsers.excel_convert import detect_excel_format, normalize_excel_bytes
from app.parsers.xlsx_merge import fill_merged_cells_xlsx
from app.parsers.xlsx_repair import repair_xlsx_bytes

logger = logging.getLogger(__name__)


class ExcelParser(BaseParser):
    """Parse every non-empty worksheet into a Markdown table."""

    def parse_into_text(self, content: bytes) -> Document:
        source_format = detect_excel_format(content) or self.file_type or "unknown"
        try:
            normalized = normalize_excel_bytes(content, file_type=self.file_type)
            repaired = repair_xlsx_bytes(normalized)
            if repaired is not None:
                normalized = repaired
            normalized = fill_merged_cells_xlsx(normalized)
            workbook = load_workbook(BytesIO(normalized), data_only=True)
        except Exception as exc:
            logger.error("Failed to open Excel workbook: %s", exc)
            return Document(content="", metadata={"error": f"open_failed: {exc}"})

        sections = []
        try:
            for sheet in workbook.worksheets:
                # TODO(security): Reject sparse or oversized worksheets before
                # calling iter_rows(). A crafted sheet can advertise Excel's
                # maximum grid size while containing very little real data,
                # causing billions of empty cells to be traversed.
                table = _sheet_to_markdown(sheet.iter_rows(values_only=True))
                if table:
                    sections.append(f"## {sheet.title}\n\n{table}")
        finally:
            workbook.close()

        return Document(
            content="\n\n".join(sections),
            metadata={
                "format": source_format,
                "sheet_count": len(workbook.sheetnames),
                "parsed_sheet_count": len(sections),
            },
        )


def _sheet_to_markdown(rows) -> str:
    # TODO(security): Stream rows and cap output size instead of materializing
    # the whole worksheet. Large inputs can exhaust memory even after dimension
    # checks if the Markdown output is allowed to grow without bounds.
    materialized = [list(row) for row in rows]
    materialized = [row for row in materialized if any(_has_value(cell) for cell in row)]
    if not materialized:
        return ""

    width = max(
        index + 1
        for row in materialized
        for index, cell in enumerate(row)
        if _has_value(cell)
    )
    normalized = [row[:width] + [None] * (width - len(row)) for row in materialized]
    headers = [
        _format_cell(value) or get_column_letter(index + 1)
        for index, value in enumerate(normalized[0])
    ]
    lines = [
        _markdown_row(headers),
        _markdown_row(["---"] * width),
    ]
    lines.extend(_markdown_row([_format_cell(value) for value in row]) for row in normalized[1:])
    return "\n".join(lines)


def _has_value(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("|", r"\|").replace("\n", "<br>").strip()


def _markdown_row(values: list[str]) -> str:
    return "| " + " | ".join(values) + " |"
