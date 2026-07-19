"""PPTX 幻灯片文本与媒体解析器。"""

import logging
import zipfile
from io import BytesIO

from pptx import Presentation

from app.parsers.base import BaseParser
from app.parsers.document import Document
from app.parsers.pptx_media import extract_pptx_media

logger = logging.getLogger(__name__)

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def normalize_pptx_bytes(content: bytes, file_type: str | None = None) -> bytes:
    """验证并返回原生 PPTX；旧 PPT 不做外部转换。"""
    normalized_type = (file_type or "").lstrip(".").lower()
    if content.startswith(_OLE_MAGIC) or normalized_type == "ppt":
        raise ValueError("legacy_ppt_not_supported")
    if not zipfile.is_zipfile(BytesIO(content)):
        raise ValueError("invalid_pptx_archive")
    with zipfile.ZipFile(BytesIO(content), "r") as archive:
        names = {name.replace("\\", "/") for name in archive.namelist()}
    if "ppt/presentation.xml" not in names:
        raise ValueError("invalid_pptx_archive: missing presentation.xml")
    return content


class PptxParser(BaseParser):
    """按幻灯片提取文本，并将内嵌媒体写入 Document.images。"""

    def parse_into_text(self, content: bytes) -> Document:
        try:
            normalized = normalize_pptx_bytes(content, self.file_type)
            presentation = Presentation(BytesIO(normalized))
            images = extract_pptx_media(normalized)
        except Exception as exc:
            logger.error("打开 PPTX 失败：%s", exc)
            return Document(content="", metadata={"error": f"open_failed: {exc}"})
        sections = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = _extract_slide_texts(slide)
            section = f"## Slide {index}"
            if texts:
                section += "\n\n" + "\n\n".join(texts)
            sections.append(section)
        if images:
            media_lines = [
                f"![{path.rsplit('/', 1)[-1]}]({path})" for path in images
            ]
            sections.append("## 媒体\n\n" + "\n".join(media_lines))
        return Document(
            content="\n\n".join(sections),
            images=images,
            metadata={
                "format": "pptx",
                "slide_count": len(presentation.slides),
                "media_count": len(images),
            },
        )


def _extract_slide_texts(slide) -> list[str]:
    """按 shape 顺序提取非空文本，表格按行拼接。"""
    texts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                texts.append(text)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" |"):
                    texts.append(text)
    return texts
