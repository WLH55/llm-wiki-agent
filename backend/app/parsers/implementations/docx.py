"""Docx2Parser：Word .docx 文档解析。

基于 python-docx 的生产基础路径：
- 用 python-docx 加载 BytesIO(content)
- 遍历 paragraphs → 段落文本
- 遍历 tables → GFM markdown 表格
- 从 OOXML package 抽取 `word/media` 图片

明确不处理：
- 并发抽取 / 多模态图像处理（docreader 用 ProcessPoolExecutor）
- 老式 .doc 二进制格式（见 doc_parser.py）
"""

import base64
import logging
from io import BytesIO
from pathlib import PurePosixPath

from docx import Document as DocxDocument

from app.parsers.core.base import BaseParser
from app.parsers.core.document import Document

logger = logging.getLogger(__name__)


class Docx2Parser(BaseParser):
    """Word .docx → 段落、GFM 表格与内嵌图片。"""

    def parse_into_text(self, content: bytes) -> Document:
        logger.info("Parsing DOCX, content size: %d bytes", len(content))

        try:
            doc = DocxDocument(BytesIO(content))
        except Exception as exc:
            logger.error("Failed to open DOCX: %s", exc)
            return Document(content="", metadata={"error": f"open_failed: {exc}"})

        parts: list[str] = []

        # 段落文本
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # 表格 → GFM markdown，首行作为表头。
        for table in doc.tables:
            rows = [
                [cell.text.strip().replace("|", "\\|") for cell in row.cells] for row in table.rows
            ]
            rows = [row for row in rows if any(row)]
            if rows:
                table_lines = [
                    "| " + " | ".join(rows[0]) + " |",
                    "| " + " | ".join("---" for _ in rows[0]) + " |",
                ]
                table_lines.extend(
                    "| " + " | ".join(row) + " |" for row in rows[1:]
                )
                parts.append("\n".join(table_lines))

        images: dict[str, str] = {}
        for part in doc.part.package.parts:
            part_name = str(part.partname).replace("\\", "/")
            if not part_name.startswith("/word/media/"):
                continue
            image_path = f"images/{PurePosixPath(part_name).name}"
            images[image_path] = base64.b64encode(part.blob).decode("ascii")
        if images:
            parts.append(
                "## 图片\n\n"
                + "\n".join(f"![{PurePosixPath(path).name}]({path})" for path in images)
            )

        text = "\n\n".join(parts)
        logger.info(
            "DOCX parsed: %d paragraphs, %d tables, %d chars output",
            len(doc.paragraphs),
            len(doc.tables),
            len(text),
        )

        return Document(
            content=text,
            images=images,
            metadata={"format": "docx", "image_count": len(images)},
        )
